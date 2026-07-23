from __future__ import annotations

import io
from dataclasses import dataclass

from docx import Document as DocxDocument
from docx.table import Table
from pypdf import PdfReader


class TextExtractionError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class TextSection:
    page_number: int | None
    text: str


def _clean_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    compact: list[str] = []
    previous_blank = False
    for line in lines:
        blank = not line
        if blank and previous_blank:
            continue
        compact.append(line)
        previous_blank = blank
    return "\n".join(compact).strip()


def _table_to_text(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [_clean_text(cell.text) for cell in row.cells]
        cells = [cell for cell in cells if cell]
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _extract_docx_text(data: bytes) -> str:
    document = DocxDocument(io.BytesIO(data))
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = _clean_text(paragraph.text)
        if text:
            blocks.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        table_text = _table_to_text(table)
        if table_text:
            blocks.append(f"Table {table_index}:\n{table_text}")

    return _clean_text("\n\n".join(blocks))


def extract_text(data: bytes, media_type: str, max_characters: int) -> list[TextSection]:
    sections: list[TextSection]

    if media_type == "application/pdf":
        try:
            reader = PdfReader(io.BytesIO(data))
            sections = [
                TextSection(page_number=index + 1, text=_clean_text(page.extract_text() or ""))
                for index, page in enumerate(reader.pages)
            ]
        except Exception as exc:  # pypdf raises several parser-specific exceptions
            raise TextExtractionError("Could not parse the PDF") from exc
    elif media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            sections = [TextSection(page_number=None, text=_extract_docx_text(data))]
        except Exception as exc:
            raise TextExtractionError("Could not parse the DOCX document") from exc
    elif media_type == "text/plain":
        try:
            text = _clean_text(data.decode("utf-8-sig"))
            sections = [TextSection(page_number=None, text=text)]
        except UnicodeDecodeError as exc:
            raise TextExtractionError("TXT files must use UTF-8 encoding") from exc
    else:
        raise TextExtractionError("Unsupported media type")

    sections = [section for section in sections if section.text]
    total_characters = sum(len(section.text) for section in sections)
    if total_characters == 0:
        raise TextExtractionError("No extractable text was found")
    if total_characters > max_characters:
        raise TextExtractionError("Extracted text exceeds the configured limit")
    return sections
