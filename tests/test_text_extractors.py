from __future__ import annotations

import io

from docx import Document

from app.services.text_extractors import extract_text


def _docx_with_table_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Main policy text")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Limit"
    table.cell(1, 0).text = "Refund"
    table.cell(1, 1).text = "14 days"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_extraction_includes_tables() -> None:
    sections = extract_text(
        _docx_with_table_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        max_characters=10_000,
    )

    text = "\n".join(section.text for section in sections)
    assert "Main policy text" in text
    assert "Name | Limit" in text
    assert "Refund | 14 days" in text
